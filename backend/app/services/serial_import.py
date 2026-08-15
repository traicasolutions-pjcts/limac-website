import csv
import hashlib
from collections.abc import Iterable
from io import BytesIO, StringIO
from typing import Any

from motor.motor_asyncio import AsyncIOMotorDatabase
from pydantic import ValidationError

from app.models.enums import ImportRowStatus
from app.repositories.products import ProductRepository
from app.schemas.imports import SerialImportResult, SerialImportRowResult, SerialImportSummary
from app.schemas.products import ProductMasterIn
from app.utils.serials import normalize_serial

REQUIRED_IMPORT_COLUMNS = {
    "serial_number",
    "product_model",
    "product_category",
    "warranty_months",
}


def checksum_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def parse_import_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(StringIO(text))
        return [dict(row) for row in reader]
    if suffix in {"xlsx", "xlsm"}:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as exc:
            raise ValueError("XLSX imports require openpyxl to be installed.") from exc
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        return [
            {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
            for row in rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    raise ValueError("Only CSV and XLSX files are supported.")


def validate_headers(rows: list[dict[str, Any]]) -> None:
    headers = set(rows[0].keys()) if rows else set()
    missing = REQUIRED_IMPORT_COLUMNS - headers
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}.")


def _clean_row(raw: dict[str, Any]) -> dict[str, Any]:
    cleaned_row: dict[str, Any] = {}
    for key, value in raw.items():
        if not key:
            continue
        if value is None:
            cleaned_row[key] = None
            continue
        cleaned = str(value).strip()
        cleaned_row[key] = cleaned or None
    return cleaned_row


def _increment(summary: SerialImportSummary, status: ImportRowStatus) -> None:
    if status == ImportRowStatus.INSERTED:
        summary.inserted += 1
    elif status == ImportRowStatus.UPDATED:
        summary.updated += 1
    elif status == ImportRowStatus.UNCHANGED:
        summary.unchanged += 1
    elif status == ImportRowStatus.SKIPPED:
        summary.skipped += 1
    elif status == ImportRowStatus.FAILED:
        summary.failed += 1


async def import_serial_rows(
    db: AsyncIOMotorDatabase,
    rows: Iterable[dict[str, Any]],
    *,
    dry_run: bool,
    checksum_sha256: str,
) -> SerialImportResult:
    repository = ProductRepository(db)
    summary = SerialImportSummary()
    results: list[SerialImportRowResult] = []

    for row_number, raw_row in enumerate(rows, start=2):
        row = _clean_row(raw_row)
        try:
            product = ProductMasterIn.model_validate(row)
        except ValidationError as exc:
            status = ImportRowStatus.FAILED
            _increment(summary, status)
            results.append(
                SerialImportRowResult(
                    row_number=row_number,
                    serial_number=row.get("serial_number"),
                    status=status,
                    reason="; ".join(error["msg"] for error in exc.errors()),
                )
            )
            continue

        status_value, reason = await repository.upsert_from_import(product, dry_run=dry_run)
        status = ImportRowStatus(status_value)
        _increment(summary, status)
        results.append(
            SerialImportRowResult(
                row_number=row_number,
                serial_number=product.serial_number,
                serial_normalized=normalize_serial(product.serial_number),
                status=status,
                reason=reason,
            )
        )

    return SerialImportResult(
        dry_run=dry_run,
        checksum_sha256=checksum_sha256,
        summary=summary,
        rows=results,
    )
